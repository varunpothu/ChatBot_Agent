from pathlib import Path
import csv,json
from knowledge.document_types import DocumentType,detect_type
from knowledge.schema import NormalizedDocument,SourceLocation

def parse_document(path:str|Path)->NormalizedDocument:
    path=Path(path); kind=detect_type(path.name)
    if kind in {DocumentType.TXT,DocumentType.MD}:
        return NormalizedDocument(path.stem,path.name,kind,path.stem,path.read_text(encoding="utf-8",errors="replace"))
    if kind==DocumentType.JSON:
        return NormalizedDocument(path.stem,path.name,kind,path.stem,json.dumps(json.loads(path.read_text(encoding="utf-8")),indent=2))
    if kind==DocumentType.CSV:
        with path.open(newline="",encoding="utf-8-sig",errors="replace") as f: rows=list(csv.reader(f))
        return NormalizedDocument(path.stem,path.name,kind,path.stem,"\n".join(" | ".join(r) for r in rows))
    if kind==DocumentType.PDF:
        from pypdf import PdfReader
        reader=PdfReader(str(path)); parts=[]; locations=[]
        for n,page in enumerate(reader.pages,1):
            text=(page.extract_text() or "").strip()
            if text: parts.append(text); locations.append(SourceLocation(page=n))
        return NormalizedDocument(path.stem,path.name,kind,path.stem,"\n\n".join(parts),tuple(locations))
    if kind==DocumentType.DOCX:
        from docx import Document
        d=Document(str(path)); text="\n".join(p.text.strip() for p in d.paragraphs if p.text.strip())
        return NormalizedDocument(path.stem,path.name,kind,path.stem,text)
    if kind==DocumentType.PPTX:
        from pptx import Presentation
        p=Presentation(str(path)); slides=[]
        for slide in p.slides:
            t="\n".join(s.text.strip() for s in slide.shapes if hasattr(s,"text") and s.text.strip())
            if t: slides.append(t)
        return NormalizedDocument(path.stem,path.name,kind,path.stem,"\n\n".join(slides))
    if kind==DocumentType.XLSX:
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True); sheets=[]
        for ws in wb.worksheets:
            rows=[" | ".join(str(v) for v in row if v is not None) for row in ws.iter_rows(values_only=True)]
            rows=[x for x in rows if x]
            if rows: sheets.append("[Sheet: "+ws.title+"]\n"+"\n".join(rows))
        return NormalizedDocument(path.stem,path.name,kind,path.stem,"\n\n".join(sheets))
    if kind==DocumentType.HTML:
        from bs4 import BeautifulSoup
        soup=BeautifulSoup(path.read_text(encoding="utf-8",errors="replace"),"html.parser")
        return NormalizedDocument(path.stem,path.name,kind,path.stem,soup.get_text("\n",strip=True))
    raise ValueError(f"Unsupported document type: {path.suffix or 'no extension'}")
